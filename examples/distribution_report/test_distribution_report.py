"""Synthetic only: no credentials or network. Run pytest in this directory."""

import importlib.util
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import UUID

import httpx
import pytest
from openpyxl import load_workbook

from asknews_sdk import AskNewsSDK


spec = importlib.util.spec_from_file_location(
    "distribution_report", Path(__file__).with_name("distribution_report.py")
)
m = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = m
spec.loader.exec_module(m)
UTC = timezone.utc


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    import socket

    monkeypatch.setattr(
        socket.socket, "connect", Mock(side_effect=AssertionError("Network forbidden"))
    )


def article(n=1, value=3, url=None):
    return {
        "article_id": str(UUID(int=n)),
        "article_url": url or f"https://apnews.com/article/{n}",
        "hit_count": value,
    }


def metadata(n=1, **overrides):
    data = {
        "article_id": str(UUID(int=n)),
        "article_url": f"https://apnews.com/article/{n}",
        "domain_url": "apnews.com",
        "classification": ["World"],
        "country": "US",
        "source_id": "synthetic",
        "page_rank": 3,
        "eng_title": f"English {n}",
        "title": f"Title {n}",
        "entities": {"Organization": ["Synthetic"]},
        "keywords": ["test"],
        "language": "en",
        "pub_date": "2026-08-01T02:00:00Z",
        "summary": "synthetic",
        "sentiment": 0,
        "key_points": ["Point"],
        "reporting_voice": "Objective",
        "provocative": "low",
    }
    data.update(overrides)
    return data


def page(rows, number=1, next_page=None):
    return {
        "data": rows,
        "total_count": sum(row["hit_count"] for row in rows),
        "page": number,
        "next_page": next_page,
    }


class Fake:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def get(self, endpoint, query):
        self.calls.append((endpoint, query))
        value = next(self.responses)
        if isinstance(value, Exception):
            raise value
        return value


@pytest.mark.parametrize(
    "month,days", [("2026-09", 30), ("2024-02", 29), ("2026-02", 28), ("2026-12", 31)]
)
def test_months_and_utc(month, days, monkeypatch):
    monkeypatch.setenv("TZ", "Pacific/Honolulu")
    start, end, following = m.month_bounds(month)
    assert following - start == timedelta(days=days)
    assert (following - end).total_seconds() == 1
    assert start.tzinfo == UTC
    if month == "2026-09":
        assert int(start.timestamp()) == 1788220800
        assert int(end.timestamp()) == 1790812799
    assert int(end.timestamp()) == int(following.timestamp()) - 1


@pytest.mark.parametrize(
    "month", ["2026-13", "2026-00", "2026-9", "1999-12", "2101-01", "2026-09Z"]
)
def test_invalid_month(month):
    with pytest.raises(m.ReportError):
        m.month_bounds(month)


@pytest.mark.parametrize(
    "value", ["https://apnews.com", "*.apnews.com", "apnews.com.evil/", "apnews.com.", "a,b.com"]
)
def test_domain_input(value):
    with pytest.raises(m.ReportError):
        m.exact_domain(value)


def test_paginate_and_page_totals():
    reader = Fake([page([article()], next_page=2), page([article(2, 4)], 2)])
    ranked = m.ranking(reader, {}, "apnews.com", "surface", 1, 4)
    assert ranked.complete and ranked.pages == 2
    assert sum(item[1] for item in ranked.rows.values()) == 7
    assert [call[1]["page"] for call in reader.calls] == [1, 2]


@pytest.mark.parametrize(
    "second",
    [
        page([article()], 2),
        page([article(2)], 2, 2),
        {"data": [], "total_count": 0},
        {"data": [article(2)], "total_count": 999, "page": 2, "next_page": None},
    ],
)
def test_page_guards_preserve_only_valid_pages(second):
    ranked = m.ranking(
        Fake([page([article()], next_page=2), second]), {}, "apnews.com", "surface", 1, 5
    )
    assert not ranked.complete
    assert len(ranked.rows) == 1


def test_caps_and_empty_final_page():
    ranked = m.ranking(Fake([page([article()], next_page=2)]), {}, "apnews.com", "surface", 1, 1)
    assert not ranked.complete and "cap" in ranked.note
    ranked = m.ranking(
        Fake([page([article()], next_page=2), page([], 2)]), {}, "apnews.com", "surface", 1, 2
    )
    assert ranked.complete and len(ranked.rows) == 1


@pytest.mark.parametrize(
    "url",
    [
        "https://other.com/a",
        "https://www.apnews.com/a",
        "https://apnews.com.evil/a",
        "https://u:p@apnews.com/a",
        '=HYPERLINK("https://evil")',
    ],
)
def test_domain_isolation(url):
    ranked = m.ranking(Fake([page([article(url=url)])]), {}, "apnews.com", "surface", 100, 1)
    assert not ranked.complete and not ranked.rows
    assert url not in ranked.note


def test_missing_id_is_not_url_joined():
    row = article()
    del row["article_id"]
    ranked = m.ranking(Fake([page([row])]), {}, "apnews.com", "surface", 100, 1)
    assert not ranked.complete and not ranked.rows


def report_fixture():
    report = m.Report("apnews.com", "2026-09")
    report.ranks = {
        "surface": m.Ranking({str(UUID(int=1)): ("https://apnews.com/a", 4)}, True, "exhausted", 1),
        "citation": m.Ranking(
            {str(UUID(int=2)): ("https://apnews.com/b", 2)}, True, "exhausted", 1
        ),
        "grounded": m.Ranking({}, False, "cap", 0),
    }
    return report


def test_join_missing_vs_zero_no_double_count():
    rows = list(m.joined_articles(report_fixture()))
    assert rows[0][2:] == (4, 0, None)
    assert rows[1][2:] == (0, 2, None)


def test_conflicting_identity_rejected():
    report = report_fixture()
    report.ranks["citation"].rows[str(UUID(int=1))] = ("https://apnews.com/changed", 2)
    with pytest.raises(m.ReportError, match="conflicting"):
        list(m.joined_articles(report))


def test_missing_daily_fields_and_absent_days_stay_unknown():
    start, _, end = m.month_bounds("2026-09")
    result = m.daily_values(
        {"data": [{"day": "2026-09-01", "surfaces": 0}], "total_surfaces": 0}, m.FIELDS, start, end
    )
    assert result["2026-09-01"]["surfaces"] == 0
    assert result["2026-09-01"]["grounded"] is None
    assert "2026-09-02" not in result


@pytest.mark.parametrize(
    "days", [["2026-09-01", "2026-09-01"], ["2026-10-01"], ["2026-09-01T00:00:00-04:00"]]
)
def test_daily_guard(days):
    start, _, end = m.month_bounds("2026-09")
    with pytest.raises(m.ReportError):
        m.daily_values(
            {"data": [{"day": day, "surfaces": 1} for day in days]}, ["surfaces"], start, end
        )


def test_daily_reconciliation():
    start, _, end = m.month_bounds("2026-09")
    with pytest.raises(m.ReportError, match="sum"):
        m.daily_values(
            {"data": [{"day": "2026-09-01", "surfaces": 3}], "total_surfaces": 4},
            ["surfaces"],
            start,
            end,
        )


def sdk_reader(handler, **kwargs):
    reader = m.Reader(None, sleep=Mock(), **kwargs)
    sdk = AskNewsSDK(
        api_key="synthetic-not-real",
        transport=httpx.MockTransport(handler),
        retries=0,
        follow_redirects=False,
        event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
    )
    reader.sdk = sdk
    return reader


@pytest.mark.parametrize("status,error", [(401, m.AuthError), (403, m.Denied)])
def test_auth_no_retries_or_body_leak(status, error):
    handler = Mock(return_value=httpx.Response(status, text="SECRET PII"))
    reader = sdk_reader(handler)
    with pytest.raises(error) as caught:
        reader.get(m.RANK, {"domain_names": ["apnews.com"]})
    assert "SECRET" not in str(caught.value)
    assert handler.call_count == 1
    reader.sdk.close()


def test_429_then_503_then_success_actual_sdk():
    handler = Mock(
        side_effect=[
            httpx.Response(429, headers={"Retry-After": "3"}),
            httpx.Response(503),
            httpx.Response(200, json={"data": []}),
        ]
    )
    reader = sdk_reader(handler)
    assert reader.get(m.RANK, {}) == {"data": []}
    assert reader.requests == 3
    assert [call.args[0] for call in reader.sleep.call_args_list] == [6.0, 6.0, 6.0]
    assert all(call.args[0].method == "GET" for call in handler.call_args_list)
    reader.sdk.close()


def test_retry_after_date_and_budget():
    now = datetime(2026, 9, 1, tzinfo=UTC)
    assert m.retry_delay("Tue, 01 Sep 2026 00:00:07 GMT", 0, now) == 7
    for value in ("600", "NaN", "infinity", "nonsense"):
        with pytest.raises(m.ReportError):
            m.retry_delay(value, 0, now)
    reader = sdk_reader(lambda req: httpx.Response(503), max_requests=1)
    with pytest.raises(m.ReportError, match="budget"):
        reader.get(m.RANK, {})
    assert reader.requests == 1
    reader.sdk.close()


def test_transient_retry_bound_and_timeout_budget():
    reader = sdk_reader(lambda req: httpx.Response(503))
    with pytest.raises(m.ReportError, match="three"):
        reader.get(m.RANK, {})
    assert reader.requests == 3
    reader.deadline = 0
    with pytest.raises(m.ReportError, match="budget"):
        reader.get(m.RANK, {})
    reader.sdk.close()


def test_oauth_401_stops_before_sdk_refresh():
    handler = Mock(return_value=httpx.Response(401, text="secret"))
    reader = m.Reader(None, sleep=Mock())
    sdk = AskNewsSDK(
        client_id="fake",
        client_secret="fake",
        scopes={"distribution"},
        transport=httpx.MockTransport(handler),
        event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
    )
    reader.sdk = sdk
    with pytest.raises(m.AuthError):
        reader.get(m.RANK, {})
    assert handler.call_count == 1  # token exchange only; no retry
    sdk.close()


def test_index_disjoint_utc_and_leap_days():
    report = m.Report("apnews.com", "2024-02")
    start, _, following = m.month_bounds(report.month)
    calls = []

    def get(endpoint, query):
        calls.append(query)
        assert endpoint == m.INDEX
        assert query["domains"] == ["apnews.com"]
        assert query["time_filter"] == "crawl_date"
        return [{"start": query["start_datetime"], "end": query["end_datetime"], "count": 1}]

    m.collect_index(SimpleNamespace(get=get), report, start, following)
    assert len(calls) == 29 and sum(report.index.values()) == 29
    for first, second in zip(calls, calls[1:]):
        assert datetime.fromisoformat(
            first["end_datetime"].replace("Z", "+00:00")
        ) < datetime.fromisoformat(second["start_datetime"].replace("Z", "+00:00"))
    assert all(query["end_datetime"].endswith("23:59:59Z") for query in calls)


@pytest.mark.parametrize("start_string", ["2026-09-01T00:00:00", "2026-09-01T00:00:00-04:00"])
def test_index_rejects_naive_or_shifted_dates(start_string):
    report = m.Report("apnews.com", "2026-09")
    start, _, end = m.month_bounds(report.month)
    with pytest.raises(m.ReportError, match="timezone"):
        m.collect_index(
            Fake([[{"start": start_string, "end": "2026-09-01T23:59:59Z", "count": 1}]]),
            report,
            start,
            end,
        )
    assert not report.index


def test_share_not_dollars_and_foreign_rejected():
    assert m.parse_share({"data": [{"domain": "apnews.com", "hit_share": 0}]}, "apnews.com") == 0
    with pytest.raises(m.ReportError):
        m.parse_share({"data": [{"domain": "other.com", "hit_share": 0.4}]}, "apnews.com")
    with pytest.raises(m.ReportError):
        m.parse_share({"data": []}, "apnews.com")


def test_internal_forbidden_group_stops_without_zero():
    report = report_fixture()
    reader = Fake([m.Denied("403: internal scope required")])
    start, _, following = m.month_bounds(report.month)
    report.section("internal", lambda: m.collect_internal(reader, report, {}, start, following))
    assert len(reader.calls) == 1 and report.issues and not report.totals


def test_workbook_open_totals_injection_and_no_links(tmp_path):
    report = report_fixture()
    report.share = 0.25
    report.daily = {"2026-09-01": {"surfaces": 0, "citations": 2, "grounded": None}}
    book = m.workbook(report)
    payloads = [
        '=HYPERLINK("https://evil","x")',
        "+SUM(A1)",
        "-1+2",
        "@SUM(A1)",
        "\t=1+1",
        "https://apnews.com/a?x==SUM(A1)",
        "title\x01control",
    ]
    m.sheet(book, "Synthetic injection", ("Text",), [(value,) for value in payloads])
    path = tmp_path / "synthetic.xlsx"
    m.save_workbook(book, path)
    opened = load_workbook(path, data_only=False)
    assert {
        "Summary",
        "Daily metrics",
        "Articles",
        "Weighted share",
        "Coverage",
        "Methodology",
    } <= set(opened.sheetnames)
    assert opened["Daily metrics"]["B2"].value == 0
    assert opened["Daily metrics"]["D2"].value is None
    assert opened["Daily metrics"]["B3"].value is None
    assert opened["Weighted share"]["B2"].value == 0.25
    assert sum(row[2] for row in opened["Articles"].iter_rows(min_row=2, values_only=True)) == 4
    assert opened["Synthetic injection"]["A2"].value == payloads[0]
    for ws in opened:
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref
        for row in ws:
            assert all(cell.data_type != "f" and not cell.hyperlink for cell in row)
    assert not opened._external_links
    with pytest.raises(m.ReportError, match="exists"):
        m.save_workbook(book, path)


def test_environment_auth_and_no_cli_credentials():
    assert m.credentials({"ASKNEWS_API_KEY": "synthetic"}) == {"api_key": "synthetic"}
    assert m.credentials({"ASKNEWS_CLIENT_ID": "a", "ASKNEWS_CLIENT_SECRET": "b"})["scopes"] == {
        "distribution",
        "news",
    }
    for env in (
        {},
        {"ASKNEWS_CLIENT_SECRET": "b"},
        {"ASKNEWS_API_KEY": "a", "ASKNEWS_CLIENT_ID": "b"},
    ):
        with pytest.raises(m.ReportError):
            m.credentials(env)


def test_end_to_end_synthetic_sdk_requests_and_metrics(tmp_path):
    def handler(request):
        q = request.url.params
        assert request.method == "GET"
        path = request.url.path
        if path == "/v1/news":
            return httpx.Response(200, json=[metadata()])
        assert q.get_list("domain_names") == ["apnews.com"] or q.get_list("domains") == [
            "apnews.com"
        ]
        if path == m.RANK:
            return httpx.Response(200, json=page([article(value=3)]))
        if path == m.SHARE:
            return httpx.Response(200, json={"data": [{"domain": "apnews.com", "hit_share": 0.1}]})
        if path == m.TOTAL:
            return httpx.Response(200, json=dict.fromkeys(m.FIELDS, 3))
        if path == m.DAILY:
            return httpx.Response(
                200,
                json={
                    "data": [dict(day="2026-09-01", **dict.fromkeys(m.FIELDS, 3))],
                    **{"total_" + name: 3 for name in m.FIELDS},
                },
            )
        if path == m.UNIQUE:
            return httpx.Response(200, json={"hits": 3, "surfaced": 1})
        if path == m.UNIQUE_DAILY:
            return httpx.Response(
                200,
                json={
                    "data": [{"day": "2026-09-01", "hits": 3, "surfaced": 1}],
                    "total_hits": 3,
                    "total_surfaced": 1,
                },
            )
        assert path == m.INDEX
        return httpx.Response(
            200, json=[{"start": q["start_datetime"], "end": q["end_datetime"], "count": 2}]
        )

    reader = sdk_reader(handler)
    report = m.collect(reader, "apnews.com", "2026-09", include_internal=True)
    assert not report.issues and reader.requests == 39
    assert report.totals["grounded"] == 3 and len(report.index) == 30
    path = tmp_path / "complete-fixture.xlsx"
    m.save_workbook(m.workbook(report), path)
    assert load_workbook(path)["Articles"].max_row == 2
    reader.sdk.close()


def test_partial_output_is_explicit(tmp_path, monkeypatch, capsys):
    report = report_fixture()
    report.issues = ["synthetic partial"]
    monkeypatch.setattr(m, "credentials", lambda env: {"api_key": "synthetic"})
    monkeypatch.setattr(m, "collect", lambda *a: report)
    output = tmp_path / "partial.xlsx"
    assert m.main(["--month", "2026-09", "--output", str(output)]) == 1
    assert not output.exists()
    assert m.main(["--month", "2026-09", "--output", str(output), "--allow-partial"]) == 2
    assert output.exists()
    assert "PARTIAL CALENDAR COVERAGE" in capsys.readouterr().out


def test_api_401_does_not_trigger_sdk_oauth_refresh():
    handler = Mock(
        side_effect=[
            httpx.Response(
                200, json={"access_token": "synthetic", "expires_in": 3600, "token_type": "Bearer"}
            ),
            httpx.Response(401, text="sensitive body"),
        ]
    )
    reader = m.Reader(None, sleep=Mock())
    with AskNewsSDK(
        client_id="fake",
        client_secret="fake",
        scopes={"distribution"},
        transport=httpx.MockTransport(handler),
        event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
    ) as sdk:
        reader.sdk = sdk
        with pytest.raises(m.AuthError):
            reader.get(m.RANK, {})
    assert handler.call_count == 2
    assert reader.requests == 2


def test_partial_page_failure_and_numeric_precision():
    ranked = m.ranking(
        Fake([page([article()], next_page=2), m.ReportError("Transient exhausted")]),
        {},
        "apnews.com",
        "surface",
        1,
        3,
    )
    assert not ranked.complete and len(ranked.rows) == 1
    for value in (None, True, -1, 2**53, 1.5, "3"):
        with pytest.raises(m.ReportError):
            m.count(value)


def test_publisher_guide_rate_limit_pacing():
    reader = m.Reader(None, sleep=Mock())
    for endpoint in (m.RANK, m.INDEX, "/oauth2/token"):
        reader.before_request(httpx.Request("GET", "https://api.asknews.app" + endpoint))
    assert [call.args[0] for call in reader.sleep.call_args_list] == [6.0, 1.0, 1.0]
    assert reader.requests == 3


def test_rate_delay_must_fit_remaining_budget():
    reader = m.Reader(None, max_seconds=5, sleep=Mock())
    with pytest.raises(m.ReportError, match="budget"):
        reader.before_request(httpx.Request("GET", "https://api.asknews.app" + m.RANK))
    assert reader.requests == 0 and reader.sleep.call_count == 0


def test_official_guide_in_workbook_methodology():
    book = m.workbook(report_fixture())
    rows = dict(book["Methodology"].iter_rows(min_row=2, values_only=True))
    assert rows["Official publisher guide"] == "https://docs.asknews.app/en/publisher"
    assert "6s" in rows["Polite requests"]
    assert "personal API keys cannot access distribution" in rows["Publisher authentication"]
    assert "batches of 100" in rows["Guide article enrichment"]
    assert "weighted" in rows["Guide/source differences"]


@pytest.mark.parametrize("header", ["120", "not-a-date", "NaN", "infinity", "-1"])
def test_unusable_retry_after_stops_all_sections(header, monkeypatch):
    clock, sent = [0.0], []
    monkeypatch.setattr(m.time, "monotonic", lambda: clock[0])

    def handler(request):
        sent.append((request.url.path, clock[0]))
        return httpx.Response(429, headers={"Retry-After": header})

    reader = sdk_reader(handler)
    reader.sleep = lambda seconds: clock.__setitem__(0, clock[0] + seconds)
    try:
        report = m.collect(reader, "apnews.com", "2026-09", include_internal=True)
        assert reader.stopped_reason and len(sent) == 1
        assert report.issues and report.share is None and not report.index
        assert all(not ranked.complete for ranked in report.ranks.values())
        # Changing sections, removing request hooks or extending time cannot reopen the reader.
        reader.deadline += 1000
        with pytest.raises(m.ReportError, match="stopped"):
            reader.get(m.INDEX, {})
        assert len(sent) == 1
    finally:
        reader.sdk.close()


def test_retry_after_final_attempt_survives_section_and_endpoint_change(monkeypatch):
    clock, sent = [0.0], []
    monkeypatch.setattr(m.time, "monotonic", lambda: clock[0])

    def handler(request):
        sent.append((request.url.path, clock[0]))
        return httpx.Response(429, headers={"Retry-After": "10"})

    reader = sdk_reader(handler)
    reader.sleep = lambda seconds: clock.__setitem__(0, clock[0] + seconds)
    try:
        report = m.collect(reader, "apnews.com", "2026-09")
        assert len(sent) == 15 and report.issues
        assert sent[0][1] == 6
        assert sent[3][1] - sent[2][1] >= 10  # final surface attempt -> citation
        assert sent[12][0] == m.INDEX  # restriction also follows shared reader into news
        assert all(b[1] - a[1] >= 10 for a, b in zip(sent, sent[1:]))
        assert reader.not_before == sent[-1][1] + 10
    finally:
        reader.sdk.close()


@pytest.mark.parametrize("budget", [15, 20])
def test_retry_after_or_wait_budget_stops_future_sends(budget, monkeypatch):
    clock, sent = [0.0], []
    monkeypatch.setattr(m.time, "monotonic", lambda: clock[0])

    def handler(request):
        sent.append(clock[0])
        return httpx.Response(429, headers={"Retry-After": "10"})

    reader = sdk_reader(handler, max_seconds=budget)
    reader.sleep = lambda seconds: clock.__setitem__(0, clock[0] + seconds)
    try:
        report = m.collect(reader, "apnews.com", "2026-09")
        assert report.issues and reader.stopped_reason
        assert sent == ([6.0] if budget == 15 else [6.0, 16.0])
    finally:
        reader.sdk.close()


def test_budget_expiring_during_wait_stops_before_send(monkeypatch):
    clock = [0.0]
    monkeypatch.setattr(m.time, "monotonic", lambda: clock[0])
    handler = Mock(return_value=httpx.Response(200, json={}))
    reader = sdk_reader(handler, max_seconds=10)
    reader.sleep = lambda seconds: clock.__setitem__(0, clock[0] + 20)
    try:
        with pytest.raises(m.ReportError, match="expired"):
            reader.get(m.RANK, {})
        assert reader.stopped_reason and handler.call_count == 0
        with pytest.raises(m.ReportError, match="stopped"):
            reader.get(m.INDEX, {})
        assert handler.call_count == 0
    finally:
        reader.sdk.close()


def distinct_fixture_handler(unique_daily):
    def handler(request):
        path, query = request.url.path, request.url.params
        if path == "/v1/news":
            payload = [metadata()]
        elif path == m.RANK:
            payload = page([article(value=3)])
        elif path == m.SHARE:
            payload = {"data": [{"domain": "apnews.com", "hit_share": 0.1}]}
        elif path == m.TOTAL:
            payload = dict.fromkeys(m.FIELDS, 3)
        elif path == m.DAILY:
            payload = {
                "data": [dict(day="2026-09-01", **dict.fromkeys(m.FIELDS, 3))],
                **{"total_" + name: 3 for name in m.FIELDS},
            }
        elif path == m.UNIQUE:
            payload = {"hits": 3, "surfaced": 1}
        elif path == m.UNIQUE_DAILY:
            payload = unique_daily
        else:
            assert path == m.INDEX
            payload = [{"start": query["start_datetime"], "end": query["end_datetime"], "count": 2}]
        return httpx.Response(200, json=payload)

    return handler


@pytest.mark.parametrize(
    "unique_daily",
    [
        {"data": [{"day": "2026-09-01", "hits": 3}], "total_hits": 3, "total_surfaced": 1},
        {"data": [{"day": "2026-09-01", "surfaced": 1}], "total_hits": 3, "total_surfaced": 1},
        {"data": [{"day": "2026-09-01", "hits": 3, "surfaced": 1}], "total_surfaced": 1},
        {"data": [{"day": "2026-09-01", "hits": 3, "surfaced": 1}], "total_hits": 3},
        {"data": [{"day": "2026-09-01", "hits": 3, "surfaced": 1}]},
        {"data": []},
        {
            "data": [{"day": "2026-09-01", "hits": 3, "surfaced": None}],
            "total_hits": 3,
            "total_surfaced": 1,
        },
        {"data": [{"day": "2026-09-01", "hits": 3}], "total_hits": 3, "total_surfaced": None},
    ],
)
def test_missing_distinct_schema_requires_partial_cli(unique_daily, tmp_path, monkeypatch):
    reader = sdk_reader(distinct_fixture_handler(unique_daily))
    try:
        report = m.collect(reader, "apnews.com", "2026-09", include_internal=True)
    finally:
        reader.sdk.close()
    assert report.issues
    assert all(row[1] != "available" for row in report.coverage if row[0] == "Internal group")
    monkeypatch.setattr(m, "collect", lambda *args: report)
    monkeypatch.setattr(m, "credentials", lambda env: {"api_key": "synthetic-not-real"})
    path = tmp_path / "requires-partial.xlsx"
    assert m.main(["--month", "2026-09", "--output", str(path)]) == 1 and not path.exists()
    assert m.main(["--month", "2026-09", "--output", str(path), "--allow-partial"]) == 2
    book = load_workbook(path)
    summary = {row[0]: row[1] for row in book["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Run status"] == "PARTIAL"
    assert summary["Monthly unique retrieved articles"] == 1
    for row in unique_daily["data"]:
        if row.get("hits") is None:
            assert book["Daily metrics"]["F2"].value is None
        if row.get("surfaced") is None:
            assert book["Daily metrics"]["G2"].value is None


def test_daily_distinct_sum_is_never_monthly_distinct():
    payload = {
        "data": [
            {"day": "2026-09-01", "hits": 2, "surfaced": 1},
            {"day": "2026-09-02", "hits": 1, "surfaced": 1},
        ],
        "total_hits": 3,
        "total_surfaced": 2,
    }
    reader = sdk_reader(distinct_fixture_handler(payload))
    try:
        report = m.collect(reader, "apnews.com", "2026-09", include_internal=True)
    finally:
        reader.sdk.close()
    assert not report.issues
    assert sum(row["surfaced"] for row in report.unique_daily.values()) == 2
    assert report.unique["surfaced"] == 1


@pytest.mark.parametrize(
    "start,stop,days,full",
    [
        ("2026-09-01T00:00:00Z", "2026-09-08T00:00:00Z", 7, False),
        ("2026-09-01T00:00:00Z", "2026-10-01T00:00:00Z", 30, True),
        ("2026-09-01T02:00:00+02:00", "2026-10-01T02:00:00+02:00", 30, True),
        ("2026-09-01T23:00:00-02:00", "2026-09-03T02:30:00Z", 2, False),
        ("2026-09-05T12:00:00Z", "2026-09-05T12:01:00Z", 1, False),
    ],
)
def test_explicit_window_normalized_utc(start, stop, days, full):
    window = m.select_window(start=start, stop=stop)
    assert window.start.tzinfo == UTC and window.stop.tzinfo == UTC
    assert window.full_month == full
    assert window.end == window.stop - timedelta(seconds=1)
    assert len(list(m.day_windows(window.start, window.stop))) == days


@pytest.mark.parametrize(
    "kwargs",
    [
        {},
        {"start": "2026-09-01T00:00:00Z"},
        {"stop": "2026-09-08T00:00:00Z"},
        {"month": "2026-09", "start": "2026-09-01T00:00:00Z"},
        {"month": "2026-09", "stop": "2026-09-08T00:00:00Z"},
        *[
            {"start": value, "stop": "2026-09-08T00:00:00Z"}
            for value in (
                "2026-09-01",
                "2026-09-01T00:00:00",
                "2026-09-01 00:00:00Z",
                "2026-09-01T00:00:00.1Z",
                "2026-09-01T00:00:00+24:00",
                "2026-09-01T00:00:00+01:60",
                "2026-09-31T00:00:00Z",
                "2026-09-08T00:00:00Z",
                "2026-09-09T00:00:00Z",
                "2026-08-31T23:59:59Z",
                "1999-12-31T00:00:00Z",
            )
        ],
        {"start": "2026-09-01T00:00:00Z", "stop": "2026-10-01T00:00:01Z"},
    ],
)
def test_invalid_range_fails_before_credentials_or_sdk(kwargs, monkeypatch, capsys):
    auth = Mock(side_effect=AssertionError("Credentials must not be read"))
    sdk = Mock(side_effect=AssertionError("SDK must not be constructed"))
    monkeypatch.setattr(m, "credentials", auth)
    monkeypatch.setattr(m, "AskNewsSDK", sdk)
    argv = [part for name, value in kwargs.items() for part in ("--" + name, value)]
    assert m.main(argv) == 1
    assert auth.call_count == sdk.call_count == 0
    assert capsys.readouterr().err


def test_week_and_partial_day_all_endpoints_and_workbook(tmp_path):
    for start, stop, days in (
        ("2026-09-01T00:00:00Z", "2026-09-08T00:00:00Z", 7),
        ("2026-09-01T12:34:56Z", "2026-09-03T06:07:08Z", 3),
    ):
        window = m.select_window(start=start, stop=stop)
        sent = []

        def handler(request, window=window, sent=sent):
            path, q = request.url.path, request.url.params
            sent.append(path)
            assert path not in (m.UNIQUE, m.UNIQUE_DAILY)
            if path in (m.RANK, m.SHARE, m.TOTAL, m.DAILY):
                assert int(q["start_date"]) == int(window.start.timestamp())
                assert int(q["end_date"]) == int(window.end.timestamp())
                assert q.get_list("domain_names") == ["apnews.com"]
            if path == m.RANK:
                payload = page([article()])
            elif path == m.SHARE:
                payload = {"data": [{"domain": "apnews.com", "hit_share": 0.2}]}
            elif path == m.TOTAL:
                payload = dict.fromkeys(m.FIELDS, 3)
            elif path == m.DAILY:
                payload = {
                    "data": [dict(day="2026-09-01", **dict.fromkeys(m.FIELDS, 3))],
                    **{"total_" + name: 3 for name in m.FIELDS},
                }
            elif path == "/v1/news":
                assert q.get_list("article_ids") == [str(UUID(int=1))]
                assert q["full_text"].lower() == "false"
                payload = [metadata()]
            else:
                assert path == m.INDEX
                begin, end = (
                    m.parse_datetime(q["start_datetime"]),
                    m.parse_datetime(q["end_datetime"]),
                )
                assert window.start <= begin <= end < window.stop
                assert begin.date() == end.date()
                payload = [{"start": q["start_datetime"], "end": q["end_datetime"], "count": 2}]
            return httpx.Response(200, json=payload)

        reader = sdk_reader(handler)
        try:
            report = m.collect(reader, "apnews.com", window, include_internal=True)
            assert not report.issues
            assert reader.requests == 3 + 1 + 2 + 1 + days
        finally:
            reader.sdk.close()
        assert not report.unique and not report.unique_daily
        assert len(report.enrichment) == 1 and len(report.index) == days
        assert any(row[1] == "unsupported for partial month" for row in report.coverage)
        path = tmp_path / f"window-{days}.xlsx"
        m.save_workbook(m.workbook(report), path)
        book = load_workbook(path)
        summary = {row[0]: row[1] for row in book["Summary"].iter_rows(min_row=2, values_only=True)}
        assert summary["Requested start UTC (inclusive)"] == start
        assert summary["Requested stop UTC (exclusive)"] == stop
        assert summary["Monthly unique retrieved articles"] is None
        assert summary["Indexed crawl-date counts (exported windows)"] == 2 * days
        assert book["Daily metrics"].max_row == days + 1
        for row in book["Daily metrics"].iter_rows(min_row=2, values_only=True):
            assert row[5] is None and row[6] is None
        assert book["Daily metrics"]["B2"].value == 3  # clipped first day is accepted
        values = list(book["Articles"].values)
        enriched = dict(zip(values[0], values[1]))
        assert enriched["pub_date"].startswith("2026-08-01")  # event cohort, not pub cohort
        assert enriched["title"] == "Title 1" and enriched["Retrieval events"] == 3


def test_named_method_batches_uuid_join_missing_enrichment_and_injection(tmp_path, monkeypatch):
    report = m.Report("apnews.com", "2026-09")
    rows = {str(UUID(int=i)): (f"https://apnews.com/article/{i}", i) for i in range(1, 202)}
    report.ranks = {metric: m.Ranking(rows.copy(), True, "exhausted") for metric in m.METRICS}
    batches = []

    def handler(request):
        assert request.url.path == "/v1/news" and request.method == "GET"
        assert request.url.params["full_text"].lower() == "false"
        ids = request.url.params.get_list("article_ids")
        batches.append(ids)
        payload = [
            metadata(UUID(value).int, title="=1+1", keywords=["@SUM(A1)"], key_points=["-1+2"])
            for value in reversed(ids)
            if UUID(value).int != 2
        ]
        return httpx.Response(200, json=payload)

    reader = sdk_reader(handler)
    named = Mock(wraps=reader.sdk.news.get_articles)
    monkeypatch.setattr(reader.sdk.news, "get_articles", named)
    try:
        report.section("Article enrichment", lambda: m.collect_enrichment(reader, report))
        assert named.call_count == reader.requests == 3
        assert [len(ids) for ids in batches] == [100, 100, 1]
        assert len({value for batch in batches for value in batch}) == 201
        for call, batch in zip(named.call_args_list, batches):
            assert call.kwargs == {"article_ids": batch, "full_text": False}
    finally:
        reader.sdk.close()
    assert len(report.enrichment) == 200 and report.issues
    path = tmp_path / "enriched.xlsx"
    m.save_workbook(m.workbook(report), path)
    book = load_workbook(path)
    values = list(book["Articles"].values)
    data = {row[0]: dict(zip(values[0], row)) for row in values[1:]}
    assert data[str(UUID(int=2))]["title"] is None
    assert data[str(UUID(int=2))]["Retrieval events"] == 2
    assert data[str(UUID(int=2))]["Enrichment coverage"] == "unavailable"
    assert data[str(UUID(int=101))]["eng_title"] == "English 101"
    assert data[str(UUID(int=101))]["Retrieval events"] == 101
    assert data[str(UUID(int=1))]["title"] == "=1+1"
    assert data[str(UUID(int=1))]["bias"] is None
    assert sum(row[2] for row in values[1:]) == sum(range(1, 202))
    for ws in book:
        for row in ws:
            assert all(cell.data_type != "f" and not cell.hyperlink for cell in row)
    assert not book._external_links


@pytest.mark.parametrize(
    "bad",
    [
        metadata(article_url="https://www.apnews.com/x"),
        metadata(domain_url="other.com"),
        metadata(3),
        metadata(article_url="https://other.com/x"),
        metadata(article_url="https://u:p@apnews.com/x"),
    ],
)
def test_enrichment_domain_and_unrequested_ids_discard_whole_batch(bad):
    report = report_fixture()
    reader = sdk_reader(lambda req: httpx.Response(200, json=[metadata(2), bad]))
    try:
        report.section("Article enrichment", lambda: m.collect_enrichment(reader, report))
    finally:
        reader.sdk.close()
    assert report.issues and not report.enrichment
    assert "other.com" not in str(report.issues)


def test_enrichment_duplicate_and_later_batch_failure_preserves_valid_data():
    report = report_fixture()
    reader = sdk_reader(lambda req: httpx.Response(200, json=[metadata(), metadata()]))
    try:
        report.section("Article enrichment", lambda: m.collect_enrichment(reader, report))
    finally:
        reader.sdk.close()
    assert not report.enrichment and report.issues
    rows = {str(UUID(int=i)): (f"https://apnews.com/article/{i}", 1) for i in range(1, 102)}
    report = m.Report("apnews.com", "2026-09")
    report.ranks = {metric: m.Ranking(rows, True) for metric in m.METRICS}

    def handler(request):
        ids = request.url.params.get_list("article_ids")
        if len(ids) == 1:
            return httpx.Response(403, text="SECRET")
        return httpx.Response(200, json=[metadata(UUID(value).int) for value in ids])

    reader = sdk_reader(handler)
    try:
        report.section("Article enrichment", lambda: m.collect_enrichment(reader, report))
        assert reader.requests == 2
    finally:
        reader.sdk.close()
    assert len(report.enrichment) == 100 and report.issues
    assert "SECRET" not in str(report.issues)


@pytest.mark.parametrize("status,error", [(401, m.AuthError), (403, m.Denied)])
def test_named_method_auth_stops_without_retry(status, error):
    reader = sdk_reader(lambda req: httpx.Response(status, text="SECRET"))
    try:
        with pytest.raises(error) as exc:
            reader.articles([str(UUID(int=1))])
        assert reader.requests == 1 and "SECRET" not in str(exc.value)
    finally:
        reader.sdk.close()


def test_named_method_shared_backpressure_and_global_caps(monkeypatch):
    clock, sent = [0.0], []
    monkeypatch.setattr(m.time, "monotonic", lambda: clock[0])

    def handler(request):
        sent.append((request.url.path, clock[0]))
        return httpx.Response(429, headers={"Retry-After": "10"})

    reader = sdk_reader(handler, max_requests=4)
    reader.sleep = lambda seconds: clock.__setitem__(0, clock[0] + seconds)
    try:
        with pytest.raises(m.ReportError, match="three"):
            reader.articles([str(UUID(int=1))])
        with pytest.raises(m.ReportError, match="budget"):
            reader.get(m.INDEX, {})
        assert [stamp for _, stamp in sent] == [1, 11, 21, 31]
        with pytest.raises(m.ReportError, match="stopped"):
            reader.articles([str(UUID(int=1))])
        assert len(sent) == 4
    finally:
        reader.sdk.close()


def test_named_oauth_token_count_and_no_refresh():
    handler = Mock(
        side_effect=[
            httpx.Response(
                200, json={"access_token": "synthetic", "expires_in": 3600, "token_type": "Bearer"}
            ),
            httpx.Response(401, text="SECRET"),
        ]
    )
    reader = m.Reader(None, sleep=Mock())
    with AskNewsSDK(
        client_id="fake",
        client_secret="fake",
        scopes={"news"},
        retries=0,
        transport=httpx.MockTransport(handler),
        event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
    ) as sdk:
        reader.sdk = sdk
        with pytest.raises(m.AuthError):
            reader.articles([str(UUID(int=1))])
    assert reader.requests == handler.call_count == 2


def test_one_second_index_clip_is_blank_without_request():
    report = report_fixture()
    window = m.select_window(start="2026-09-01T23:59:59Z", stop="2026-09-02T00:00:00Z")
    reader = Fake([])
    m.collect_index(reader, report, window.start, window.stop)
    assert not reader.calls and not report.index and report.issues


def test_empty_cohort_never_requests_enrichment():
    report = report_fixture()
    report.ranks = {metric: m.Ranking({}, True) for metric in m.METRICS}
    reader = SimpleNamespace(articles=Mock(side_effect=AssertionError("No batch expected")))
    assert m.collect_enrichment(reader, report)
    reader.articles.assert_not_called()


def test_index_partial_days_exact_clips():
    window = m.select_window(start="2026-09-01T14:34:56+02:00", stop="2026-09-03T06:07:08Z")
    report, calls = report_fixture(), []

    def get(endpoint, query):
        calls.append((query["start_datetime"], query["end_datetime"]))
        return [{"start": query["start_datetime"], "end": query["end_datetime"], "count": 0}]

    m.collect_index(SimpleNamespace(get=get), report, window.start, window.stop)
    assert calls == [
        ("2026-09-01T12:34:56Z", "2026-09-01T23:59:59Z"),
        ("2026-09-02T00:00:00Z", "2026-09-02T23:59:59Z"),
        ("2026-09-03T00:00:00Z", "2026-09-03T06:07:07Z"),
    ]
    assert list(report.index.values()) == [0, 0, 0]


def test_explicit_month_equals_shorthand():
    assert m.select_window(month="2026-09") == m.select_window(
        start="2026-09-01T00:00:00Z", stop="2026-10-01T00:00:00Z"
    )


def test_enrichment_missing_schema_sanitized_partial_and_cli_gate(tmp_path, monkeypatch):
    report = report_fixture()
    reader = sdk_reader(lambda req: httpx.Response(200, json=[{"title": "SECRET"}]))
    try:
        report.section("Article enrichment", lambda: m.collect_enrichment(reader, report))
    finally:
        reader.sdk.close()
    assert report.issues and not report.enrichment and "SECRET" not in str(report.issues)
    monkeypatch.setattr(m, "credentials", lambda env: {"api_key": "synthetic"})
    captured = []

    def collect(reader, domain, window, *args):
        captured.append(window)
        report.window = window
        return report

    monkeypatch.setattr(m, "collect", collect)
    output = tmp_path / "missing-enrichment.xlsx"
    argv = [
        "--start",
        "2026-09-01T02:00:00+02:00",
        "--stop",
        "2026-09-08T00:00:00Z",
        "--output",
        str(output),
    ]
    assert m.main(argv) == 1 and not output.exists()
    assert m.main([*argv, "--allow-partial"]) == 2
    assert captured[0].start == datetime(2026, 9, 1, tzinfo=UTC)
    assert captured[0].stop == datetime(2026, 9, 8, tzinfo=UTC)
    book = load_workbook(output)
    assert book["Articles"]["F2"].value is None
    assert book["Articles"]["C2"].value == 4


@pytest.mark.parametrize("header", ["120", "not-a-date"])
def test_named_unusable_backpressure_stops_every_send(header):
    handler = Mock(return_value=httpx.Response(429, headers={"Retry-After": header}))
    reader = sdk_reader(handler)
    try:
        with pytest.raises(m.ReportError, match="stopped"):
            reader.articles([str(UUID(int=1))])
        with pytest.raises(m.ReportError, match="stopped"):
            reader.get(m.RANK, {})
        assert reader.requests == handler.call_count == 1
    finally:
        reader.sdk.close()


def test_named_time_and_oauth_request_caps():
    handler = Mock(return_value=httpx.Response(200, json=[]))
    reader = sdk_reader(handler, max_seconds=1)
    try:
        with pytest.raises(m.ReportError, match="budget"):
            reader.articles([str(UUID(int=1))])
        assert handler.call_count == reader.requests == 0
    finally:
        reader.sdk.close()
    handler = Mock(
        return_value=httpx.Response(
            200, json={"access_token": "synthetic", "expires_in": 3600, "token_type": "Bearer"}
        )
    )
    reader = m.Reader(None, max_requests=1, sleep=Mock())
    with AskNewsSDK(
        client_id="fake",
        client_secret="fake",
        scopes={"news"},
        retries=0,
        transport=httpx.MockTransport(handler),
        event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
    ) as sdk:
        reader.sdk = sdk
        with pytest.raises(m.ReportError, match="budget"):
            reader.articles([str(UUID(int=1))])
        assert handler.call_count == reader.requests == 1  # token only; no unbudgeted news send
