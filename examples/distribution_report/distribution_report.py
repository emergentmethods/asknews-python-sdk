#!/usr/bin/env python3
"""Bounded, read-only publisher Excel export. See README.md; Python 3.10+."""

from __future__ import annotations

import argparse
import calendar
import json
import logging
import math
import os
import re
import sys
import tempfile
import time
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import urlsplit
from uuid import UUID

import httpx
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from asknews_sdk import AskNewsSDK
from asknews_sdk.errors import APIError


UTC = timezone.utc
METRICS = ("surface", "citation", "grounded")
FIELDS = ("surfaces", "citations", "grounded", "full_text")
PUBLISHER_GUIDE = "https://docs.asknews.app/en/publisher"
API_REFERENCE = "https://docs.asknews.app/en/reference#tag--distribution"
BASE = "/v1/distribution/"
RANK = BASE + "articles/top_n_for_domains"
SHARE = BASE + "stats/hit_share"
TOTAL = BASE + "stats/metrics"
DAILY = BASE + "stats/metrics_timeseries"
UNIQUE = BASE + "articles/domain_hits_surface"
UNIQUE_DAILY = BASE + "articles/domain_hits_surface_timewindow"
INDEX = "/v1/index_counts"
# Narrow transport fallbacks only: no ranking/share/unique helpers; internal DTOs
# discard grounded; get_index_counts strips UTC offsets. Recipe 4 uses a named helper.
ALLOWED = {RANK, SHARE, TOTAL, DAILY, UNIQUE, UNIQUE_DAILY, INDEX}
ENRICH_FIELDS = (
    "title",
    "eng_title",
    "pub_date",
    "language",
    "country",
    "classification",
    "sentiment",
    "entities",
    "keywords",
    "reporting_voice",
    "provocative",
    "page_rank",
    "key_points",
    "bias",
    "content_type",
)
BOUNDARY = (
    "PARTIAL CALENDAR COVERAGE: integer-second inclusive API end = exclusive stop minus 1s; "
    "events after that instant and before stop are omitted. Index repeats this gap at "
    "each clipped UTC-day stop. Never an exact full-range/full-month claim."
)


class ReportError(Exception):
    """Sanitized, operator-actionable message only; never include response bodies."""


class AuthError(ReportError):
    pass


class Denied(ReportError):
    pass


def month_bounds(month):
    if not re.fullmatch(r"20\d{2}-\d{2}|2100-\d{2}", month):
        raise ReportError("Use YYYY-MM within 2000–2100.")
    year, number = map(int, month.split("-"))
    try:
        start = datetime(year, number, 1, tzinfo=UTC)
        following = start + timedelta(days=calendar.monthrange(year, number)[1])
    except ValueError:
        raise ReportError("Invalid calendar month.") from None
    return start, following - timedelta(seconds=1), following


@dataclass(frozen=True)
class Window:
    start: datetime
    stop: datetime

    def __post_init__(self):
        for value in (self.start, self.stop):
            if value.tzinfo is None or value.utcoffset() is None or value.microsecond:
                raise ReportError("Use timezone-aware datetimes at whole-second precision.")
        object.__setattr__(self, "start", self.start.astimezone(UTC))
        object.__setattr__(self, "stop", self.stop.astimezone(UTC))
        if self.stop <= self.start:
            raise ReportError("--stop must be later than --start (exclusive stop).")
        if not 2000 <= self.start.year <= 2100:
            raise ReportError("Range must be within 2000–2100.")
        if (self.start.year, self.start.month) != (self.end.year, self.end.month):
            raise ReportError("Cross-month UTC ranges are unsupported; select one month at a time.")

    @property
    def end(self):
        return self.stop - timedelta(seconds=1)

    @property
    def full_month(self):
        start, _, stop = month_bounds(self.start.strftime("%Y-%m"))
        return (self.start, self.stop) == (start, stop)


def parse_datetime(value):
    # Explicit calendar date, seconds and numeric offset/Z only, not date-only/naive input.
    if not re.fullmatch(
        r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?(?:Z|[+-]\d{2}:\d{2})", value
    ):
        raise ReportError("Use ISO8601 YYYY-MM-DDTHH:MM:SSZ or an explicit ±HH:MM offset.")
    # datetime.fromisoformat normalizes malformed offsets such as +01:60; reject them.
    if value[-1] != "Z" and (int(value[-5:-3]) > 23 or int(value[-2:]) > 59):
        raise ReportError("Invalid ISO8601 timezone offset.")
    try:
        date = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        raise ReportError("Invalid ISO8601 calendar datetime.") from None
    if date.microsecond:
        raise ReportError("Fractional-second inputs are unsupported; use whole seconds explicitly.")
    return date.astimezone(UTC)


def select_window(month=None, start=None, stop=None):
    if month is not None:
        if start is not None or stop is not None:
            raise ReportError("--month is mutually exclusive with --start/--stop.")
        first, _, following = month_bounds(month)
        return Window(first, following)
    if start is None or stop is None:
        raise ReportError("Supply both --start and --stop, or the optional --month shorthand.")
    return Window(parse_datetime(start), parse_datetime(stop))


def day_windows(start, stop):
    current = start
    while current < stop:
        following = min(current.replace(hour=0, minute=0, second=0) + timedelta(days=1), stop)
        yield current, following
        current = following


def exact_domain(value):
    value = value.lower()
    if len(value) > 253 or not re.fullmatch(
        r"(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+"
        r"[a-z](?:[a-z0-9-]{0,61}[a-z0-9])?",
        value,
    ):
        raise ReportError("Use one bare exact DNS domain, without URL, wildcard or trailing dot.")
    return value


def count(value):
    if type(value) is not int or not 0 <= value <= 2**53 - 1:
        raise ReportError("Missing/invalid count or count exceeds exact Excel numeric precision.")
    return value


def utc_string(value):
    return value.astimezone(UTC).isoformat().replace("+00:00", "Z")


def retry_delay(header, attempt, now=None):
    delay = 2**attempt
    if header is not None:
        try:
            delay = float(header)
            if delay < 0:
                raise ReportError("Invalid negative Retry-After; stop and retry later.")
        except (TypeError, ValueError):
            try:
                date = parsedate_to_datetime(header)
                delay = (date - (now or datetime.now(UTC))).total_seconds()
            except (TypeError, ValueError, OverflowError):
                raise ReportError("Invalid Retry-After; stop and retry later.") from None
    if not math.isfinite(delay) or delay > 60:
        raise ReportError("Retry-After exceeds 60-second wait budget; retry later.")
    return max(0, delay)


class Reader:
    """Shared safety for named SDK methods and narrow transport fallbacks, including OAuth."""

    def __init__(self, sdk, max_requests=400, max_seconds=2400, sleep=time.sleep):
        self.sdk = sdk
        self.max_requests = max_requests
        self.deadline = time.monotonic() + max_seconds
        self.requests = 0
        self.sleep = sleep
        self.not_before = 0.0
        self.stopped_reason = None

    def stop(self, reason):
        self.stopped_reason = reason
        raise ReportError("All further sends stopped: " + reason)

    def defer(self, delay):
        # Shared across every section, including after a section's final failed attempt.
        self.not_before = max(self.not_before, time.monotonic() + delay)
        if self.not_before >= self.deadline:
            self.stop("Retry-After/backoff would exceed time budget.")

    def before_request(self, request):
        if self.stopped_reason is not None:
            self.stop(self.stopped_reason)
        # Publisher guide Recipe 3: 6s keeps distribution below its 0.2 req/s limit.
        # Apply even to admin/internal calls; do not assume rate-limit exemptions.
        pacing = 6.0 if request.url.path.startswith(BASE) else 1.0
        now = time.monotonic()
        delay = max(pacing, self.not_before - now)
        if self.requests >= self.max_requests or now + delay >= self.deadline:
            self.stop("Request/time budget reached; increase bounded limits explicitly.")
        # Called by httpx for API and SDK-managed OAuth requests, not logged.
        self.sleep(delay)
        if time.monotonic() >= self.deadline:
            self.stop("Time budget expired while waiting.")
        self.requests += 1

    def after_response(self, response):
        # Stop before the SDK's built-in OAuth 401 refresh, including non-JSON errors.
        if response.status_code == 401:
            raise AuthError("401: verify existing authorized credentials; no retry.")
        if response.status_code == 403:
            raise Denied(
                "403: required scope/domain authorization missing; admin is not internal scope. "
                "Use existing authorized credentials; no access changed and no retry."
            )
        # Surface even HTML/proxy errors before SDK JSON error parsing.
        response.raise_for_status()

    def get(self, endpoint, query):
        if endpoint not in ALLOWED:
            raise ReportError("Endpoint not allowlisted.")
        return self.read(
            lambda: self.sdk.client.request(method="GET", endpoint=endpoint, query=query).content
        )

    def articles(self, ids):
        # Real published SDK helper, not a local imitation; no full-text body fetch.
        return self.read(lambda: self.sdk.news.get_articles(article_ids=ids, full_text=False))

    def read(self, operation):
        for attempt in range(3):
            if self.stopped_reason is not None:
                self.stop(self.stopped_reason)
            try:
                return operation()
            except (APIError, httpx.HTTPStatusError) as exc:
                status = exc.response.status_code
                if status == 401:
                    raise AuthError(
                        "401: verify existing authorized credentials; no retry."
                    ) from None
                if status == 403:
                    raise Denied(
                        "403: credential lacks required scope/domain authorization; admin is not "
                        "internal scope. Use an already-authorized credential; no access changed."
                    ) from None
                if status not in (408, 429, 500, 502, 503, 504):
                    raise ReportError(
                        f"HTTP {status}: section unavailable; check API compatibility."
                    ) from None
                try:
                    delay = retry_delay(exc.response.headers.get("retry-after"), attempt)
                except ReportError as error:
                    # A section-local failure cannot release a shared-quota restriction.
                    self.stop(str(error))
            except httpx.TransportError:
                delay = 2**attempt
            except (ValueError, TypeError, AttributeError):
                raise ReportError("Unrecognized SDK/API response; verify compatibility.") from None
            self.defer(delay)
            if attempt == 2:
                raise ReportError("Transient request failed after three bounded attempts.")
        raise ReportError("Request did not complete.")


@dataclass
class Ranking:
    rows: dict = field(default_factory=dict)
    complete: bool = False
    note: str = "not requested"
    pages: int = 0


def parse_article(row, domain):
    url = row.get("article_url")
    if not isinstance(url, str) or len(url) > 8192:
        raise ReportError("Invalid article URL.")
    parts = urlsplit(url)
    if (
        parts.scheme not in ("http", "https")
        or parts.hostname != domain
        or parts.username
        or parts.password
        or parts.port not in (None, 80, 443)
    ):
        raise ReportError("Exact-domain isolation failed; page discarded (no foreign data saved).")
    # IDs are stable join keys; refuse URL-only joins because different IDs can share a URL.
    try:
        identity = str(UUID(row["article_id"]))
    except (KeyError, ValueError, TypeError, AttributeError):
        raise ReportError(
            "Stable article_id missing/invalid; unsafe metric join refused."
        ) from None
    return identity, (url, count(row.get("hit_count")))


def ranking(reader, query, domain, metric, page_size, max_pages):
    result = Ranking()
    try:
        for page in range(1, max_pages + 1):
            data = reader.get(RANK, dict(query, metric=metric, limit=page_size, page=page))
            rows = data["data"]
            if not isinstance(rows, list) or len(rows) > page_size:
                raise ReportError("Invalid ranking page size.")
            if type(data.get("page")) is not int or data["page"] != page or "next_page" not in data:
                raise ReportError(
                    "Pagination metadata missing/mismatched; legacy top-N may be capped."
                )
            parsed = {}
            for row in rows:
                key, item = parse_article(row, domain)
                if key in parsed or key in result.rows:
                    raise ReportError(
                        "Repeated article across pages; snapshot drift/page loop detected."
                    )
                parsed[key] = item
            if count(data.get("total_count")) != sum(item[1] for item in parsed.values()):
                raise ReportError("Page-only total_count does not reconcile.")
            next_page = data["next_page"]
            if next_page is not None and (
                type(next_page) is not int or next_page != page + 1 or not rows
            ):
                raise ReportError("Non-forward/empty pagination loop refused.")
            result.rows.update(parsed)
            result.pages += 1
            if next_page is None:
                result.complete = True
                result.note = "pagination exhausted; snapshot and range-end caveats still apply"
                return result
        raise ReportError("Page cap reached; exported counts are partial, not domain totals.")
    except (AuthError, Denied):
        raise
    except ReportError as exc:
        result.note = str(exc)
    except (KeyError, TypeError, ValueError, AttributeError):
        result.note = "Unexpected ranking schema; section partial/unavailable."
    return result


def daily_values(payload, fields, start, following, total_prefix="total_"):
    rows = payload["data"]
    if not isinstance(rows, list) or len(rows) > 31:
        raise ReportError("Invalid daily rows.")
    parsed = {}
    for row in rows:
        day = row["day"]
        if not isinstance(day, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
            raise ReportError("Daily bucket must be a UTC calendar date.")
        date = datetime.strptime(day, "%Y-%m-%d").replace(tzinfo=UTC)
        if not start.replace(hour=0, minute=0, second=0) <= date < following or day in parsed:
            raise ReportError("Duplicate/out-of-window daily bucket.")
        parsed[day] = {name: count(row[name]) if name in row else None for name in fields}
    for name in fields:
        if any(row[name] is None for row in parsed.values()) or total_prefix + name not in payload:
            continue  # Missing additive fields stay unavailable, never invented zeros.
        if count(payload[total_prefix + name]) != sum(row[name] for row in parsed.values()):
            raise ReportError("Daily sum does not match response total.")
    return parsed


@dataclass
class Report:
    domain: str
    month: str
    window: Window | None = None
    enrichment: dict = field(default_factory=dict)
    started: str = field(default_factory=lambda: utc_string(datetime.now(UTC)))
    finished: str = ""
    ranks: dict = field(default_factory=dict)
    totals: dict = field(default_factory=dict)
    daily: dict = field(default_factory=dict)
    unique: dict = field(default_factory=dict)
    unique_daily: dict = field(default_factory=dict)
    index: dict = field(default_factory=dict)
    share: float | None = None
    coverage: list = field(default_factory=list)
    issues: list = field(default_factory=list)

    @property
    def selected(self):
        return self.window or select_window(month=self.month)

    def section(self, name, operation):
        try:
            issues_before = len(self.issues)
            value = operation()
            status = "partial" if len(self.issues) > issues_before else "available"
            self.coverage.append((name, status, "See coverage issues and boundary methodology"))
            return value
        except AuthError:
            raise
        except ReportError as exc:
            message = str(exc)
        except (KeyError, TypeError, ValueError, AttributeError):
            message = "Unexpected schema; section unavailable."
        self.coverage.append((name, "unavailable", message))
        self.issues.append(name + ": " + message)
        return None


def parse_share(payload, domain):
    rows = payload["data"]
    if not isinstance(rows, list) or len(rows) != 1 or rows[0].get("domain") != domain:
        raise ReportError("Missing/non-exact share row; no other publisher data saved.")
    value = rows[0]["hit_share"]
    if type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= 1:
        raise ReportError("Invalid weighted share.")
    return value


def validate_distinct_coverage(report, payload):
    for name in ("hits", "surfaced"):
        if "total_" + name not in payload:
            report.issues.append("Daily distinct response total missing: total_" + name)
        else:
            # Validate totals even when missing row fields prevent sum reconciliation.
            count(payload["total_" + name])
        if any(row[name] is None for row in report.unique_daily.values()):
            report.issues.append("Daily distinct field missing: " + name)


def collect_internal(reader, report, query, start, following):
    # One denied internal request stops the internal group; no 403 endpoint-probing loop.
    raw = reader.get(TOTAL, query)
    report.totals = {name: count(raw[name]) if name in raw else None for name in FIELDS}
    if any(value is None for value in report.totals.values()):
        report.issues.append("Internal totals: missing fields remain unavailable.")
    report.coverage.append(("Internal totals", "available", "Raw event counts, not weighted"))
    daily_payload = reader.get(DAILY, query)
    report.daily = daily_values(daily_payload, FIELDS, start, following)
    if any("total_" + name not in daily_payload for name in FIELDS):
        report.issues.append("Daily response totals missing; completeness not established.")
    for name in FIELDS:
        values = [row[name] for row in report.daily.values()]
        if any(value is None for value in values):
            report.issues.append("Daily metric field missing: " + name)
        elif report.totals[name] is not None and sum(values) != report.totals[name]:
            report.issues.append("Daily/aggregate mismatch: " + name)
    if Window(start, following).full_month:
        collect_distinct(reader, report, start, following)
    else:
        report.coverage.append(
            (
                "Distinct retrieval counts",
                "unsupported for partial month",
                "Both distinct routes accept year/month only; not requested. Monthly and daily "
                "distinct cells remain blank. Never sum daily uniques into range uniques.",
            )
        )
    for metric, name in zip(METRICS, FIELDS):
        ranked = report.ranks[metric]
        if ranked.complete and report.totals[name] is not None:
            if sum(row[1] for row in ranked.rows.values()) != report.totals[name]:
                report.issues.append(f"{metric}: ranking/raw aggregate mismatch (snapshot drift).")
    return True


def collect_distinct(reader, report, start, following):
    # These routes have year/month ONLY, not arbitrary date bounds.
    month_query = {"domain_names": [report.domain], "year": start.year, "month": start.month}
    raw = reader.get(UNIQUE, month_query)
    report.unique = {name: count(raw[name]) for name in ("hits", "surfaced")}
    unique_daily_payload = reader.get(UNIQUE_DAILY, month_query)
    report.unique_daily = daily_values(unique_daily_payload, ("hits", "surfaced"), start, following)
    validate_distinct_coverage(report, unique_daily_payload)
    daily_hits = [row["hits"] for row in report.unique_daily.values()]
    if all(value is not None for value in daily_hits):
        if sum(daily_hits) != report.unique["hits"]:
            report.issues.append("Surface totals/daily mismatch (snapshot drift).")
    if report.unique["hits"] != report.totals["surfaces"]:
        report.issues.append("Surface totals/aggregate mismatch (snapshot drift).")
    # total_surfaced reconciles daily rows only; NEVER replace the monthly distinct count.
    surface = report.ranks["surface"]
    if surface.complete and len(surface.rows) != report.unique["surfaced"]:
        report.issues.append("Monthly distinct/ranking article count mismatch (snapshot drift).")


def collect_index(reader, report, start, following):
    for day, stop in day_windows(start, following):
        end = stop - timedelta(seconds=1)
        if end == day:
            # Server create_sub_ranges emits no bucket for equal endpoints.
            report.issues.append("Index: one-second clipped window unsupported; left blank.")
            continue
        # UTC-day windows clipped to selected bounds; no adjacent inclusive overlap.
        # Transport intentionally preserves UTC offsets lost by the named index helper.
        payload = reader.get(
            INDEX,
            {
                "domains": [report.domain],
                "start_datetime": utc_string(day),
                "end_datetime": utc_string(end),
                "sampling": "1d",
                "time_filter": "crawl_date",
            },
        )
        if not isinstance(payload, list) or len(payload) != 1:
            raise ReportError("Expected one index bucket; missing is not zero.")
        row = payload[0]
        for name, expected in (("start", day), ("end", end)):
            date = datetime.fromisoformat(row[name].replace("Z", "+00:00"))
            if date.tzinfo is None or date.astimezone(UTC) != expected:
                raise ReportError("Index bucket timezone/boundaries differ; bucket discarded.")
        report.index[day.date().isoformat()] = count(row["count"])
    return True


def collect(reader, domain, month, page_size=100, max_pages=100, include_internal=False):
    window = month if isinstance(month, Window) else select_window(month=month)
    start, end, following = window.start, window.end, window.stop
    report = Report(domain, start.strftime("%Y-%m"), window=window)
    query = {
        "domain_names": [domain],
        "start_date": int(start.timestamp()),
        "end_date": int(end.timestamp()),
    }
    for metric in METRICS:
        result = ranking(reader, query, domain, metric, page_size, max_pages)
        report.ranks[metric] = result
        report.coverage.append(
            (metric + " ranking", "exhausted" if result.complete else "partial", result.note)
        )
        if not result.complete:
            report.issues.append(metric + ": " + result.note)
    report.share = report.section(
        "Weighted share",
        lambda: parse_share(reader.get(SHARE, dict(query, is_publisher=True)), domain),
    )
    if include_internal:
        report.section(
            "Internal group", lambda: collect_internal(reader, report, query, start, following)
        )
    else:
        report.coverage.append(
            (
                "Internal group",
                "not requested",
                "Requires existing internal scope; use --include-internal",
            )
        )
    report.section("Article enrichment", lambda: collect_enrichment(reader, report))
    report.section("Index counts", lambda: collect_index(reader, report, start, following))
    report.finished = utc_string(datetime.now(UTC))
    return report


def joined_articles(report):
    joined = {}
    for metric, ranked in report.ranks.items():
        for identity, (url, value) in ranked.rows.items():
            row = joined.setdefault(identity, {"url": url})
            if row["url"] != url:
                raise ReportError("Article ID has conflicting URLs across metrics; join refused.")
            row[metric] = value
    for identity, row in sorted(joined.items()):
        values = []
        for metric in METRICS:
            # Absence means zero only for exhausted positive-event rankings.
            values.append(row.get(metric, 0 if report.ranks[metric].complete else None))
        yield (identity, row["url"], *values)


def collect_enrichment(reader, report):
    # Union ONLY the validated range-event cohorts; publication date need not be in range.
    ids = [row[0] for row in joined_articles(report)]
    for offset in range(0, len(ids), 100):
        batch = ids[offset : offset + 100]
        articles = reader.articles(batch)
        if not isinstance(articles, list) or len(articles) > len(batch):
            raise ReportError("Invalid enrichment batch size; batch discarded.")
        parsed = {}
        for article in articles:
            data = article.model_dump(mode="json", exclude_unset=True)
            identity, _ = parse_article(dict(data, hit_count=0), report.domain)
            if (
                identity not in batch
                or identity in parsed
                or data.get("domain_url") != report.domain
            ):
                raise ReportError("Enrichment ID/domain isolation failed; batch discarded.")
            parsed[identity] = {name: data.get(name) for name in ENRICH_FIELDS}
        report.enrichment.update(parsed)  # only publish after validating the entire batch
        if len(parsed) != len(batch):
            report.issues.append(
                f"Article enrichment: {len(batch) - len(parsed)} requested IDs absent in batch; "
                "retained-index/access gaps remain blank, not zero."
            )
    return True


def enriched_rows(report):
    for row in joined_articles(report):
        data = report.enrichment.get(row[0], {})
        values = []
        for name in ENRICH_FIELDS:
            value = data.get(name)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            values.append(value)
        yield (*row, *values, "returned" if row[0] in report.enrichment else "unavailable")


def sheet(workbook, title, headers, rows):
    ws = workbook.create_sheet(title)
    for row in (headers, *rows):
        ws.append(
            [
                ILLEGAL_CHARACTERS_RE.sub("", value)[:32767] if isinstance(value, str) else value
                for value in row
            ]
        )
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)[:32767]
                cell.data_type = "s"  # Explicit literal even for =,+,-,@ and URLs; no hyperlinks.
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = min(
            70,
            max(
                18,
                max(len(str(ws.cell(r, index).value or "")) for r in range(1, ws.max_row + 1)) + 2,
            ),
        )
    return ws


def workbook(report):
    window = report.selected
    start, end, following = window.start, window.end, window.stop
    book = Workbook()
    book.remove(book.active)
    summary = [
        ("Domain (exact)", report.domain, "No subdomains or other publishers"),
        ("Requested start UTC (inclusive)", utc_string(start), BOUNDARY),
        ("Requested stop UTC (exclusive)", utc_string(following), "Never expanded to month"),
        (
            "Run status",
            "PARTIAL" if report.issues else "COLLECTED WITH COVERAGE CAVEATS",
            "Blank = unavailable/not requested; zero = observed or documented absence",
        ),
        (
            "Monthly unique retrieved articles",
            report.unique.get("surfaced"),
            "Full-month requests only; partial-month unsupported/blank; never sum daily uniques",
        ),
        (
            "Exported surface article IDs",
            len(report.ranks["surface"].rows),
            "Exported cohort only; not indexed article inventory",
        ),
        (
            "Indexed crawl-date counts (exported windows)",
            sum(report.index.values()) if report.index else None,
            f"{len(report.index)} clipped daily buckets; fractional-second gaps each window; "
            "not publication/event dates",
        ),
    ]
    for metric, name in zip(METRICS, FIELDS):
        rank = report.ranks[metric]
        summary.extend(
            [
                (
                    name + " raw aggregate",
                    report.totals.get(name),
                    "Internal endpoint, if authorized",
                ),
                (
                    metric + " exported article sum",
                    sum(row[1] for row in rank.rows.values())
                    if rank.complete or rank.rows
                    else None,
                    rank.note,
                ),
            ]
        )
    summary.append(
        (
            "full_text raw aggregate",
            report.totals.get("full_text"),
            "Separate metric, not retrievals or groundings",
        )
    )
    sheet(book, "Summary", ("Metric", "Value", "Interpretation"), summary)
    daily = []
    for day, stop in day_windows(start, following):
        key = day.date().isoformat()
        counts = report.daily.get(key, {})
        unique = report.unique_daily.get(key, {})
        daily.append(
            (
                key,
                *(counts.get(name) for name in FIELDS),
                unique.get("hits"),
                unique.get("surfaced"),
                report.index.get(key),
                "Missing buckets/fields blank, not zero; "
                + utc_string(day)
                + " <= requested time < "
                + utc_string(stop),
            )
        )
    sheet(
        book,
        "Daily metrics",
        (
            "UTC day",
            *FIELDS,
            "Monthly-series hits",
            "Daily distinct retrieved",
            "Indexed crawl-date",
            "Coverage",
        ),
        daily,
    )
    sheet(
        book,
        "Articles",
        (
            "Article ID",
            "URL (literal)",
            "Retrieval events",
            "Citation events",
            "Grounding events",
            *ENRICH_FIELDS,
            "Enrichment coverage",
        ),
        list(enriched_rows(report)),
    )
    ws = sheet(
        book,
        "Weighted share",
        ("Exact domain", "Weighted fraction", "Meaning"),
        [
            (
                report.domain,
                report.share,
                "Publisher-only weighted pool fraction; metric weights/multipliers "
                "are server-configured, not returned. NOT payout or dollars.",
            )
        ],
    )
    ws["B2"].number_format = "0.000000%"
    sheet(
        book,
        "Coverage",
        ("Section", "Status", "Details"),
        report.coverage + [("Issue", "partial", issue) for issue in report.issues],
    )
    sheet(
        book,
        "Methodology",
        ("Item", "Value"),
        [
            ("Official publisher guide", PUBLISHER_GUIDE),
            ("Linked API reference", API_REFERENCE),
            (
                "Publisher authentication",
                "Guide uses existing publisher organization API keys (ank_org_...) with "
                "distribution/news scopes; personal API keys cannot access distribution. "
                "OAuth is an SDK-supported admin alternative only when already authorized. "
                "Admin status never implies internal scope.",
            ),
            (
                "Guide/source differences",
                "Guide describes traffic share broadly; implementation is weighted, not raw "
                "traffic or payout. Guide omits metric selector and index counts; public API "
                "reference/source support them. Named SDK internal DTOs still drop grounded.",
            ),
            (
                "Guide article enrichment",
                "Recipe 4: this report calls sdk.news.get_articles "
                "(article_ids=..., full_text=False) in "
                "batches of 100, joined by UUID with exact-domain validation. "
                "Only IDs from selected "
                "range-event rankings are enriched; older publication dates are legitimate. "
                "Missing metadata/retained-index gaps stay blank. No article body fetched.",
            ),
            (
                "Polite requests",
                "Guide limit: distribution 0.2 requests/sec, burst 5, concurrency 2. "
                "This report is serial and waits 6s before every distribution send, including "
                "retries; other/auth sends wait 1s. A shared Retry-After deadline applies "
                "across sections, including final attempts. Unusable headers or waits/budgets "
                "that prevent compliance permanently stop all further sends for this run.",
            ),
            ("Requested start UTC", utc_string(start)),
            ("Requested exclusive end UTC", utc_string(following)),
            ("API inclusive start timestamp", int(start.timestamp())),
            ("API inclusive end timestamp", int(end.timestamp())),
            ("Calendar coverage", BOUNDARY),
            ("Snapshot started UTC", report.started),
            ("Snapshot finished UTC", report.finished),
            (
                "Consistency",
                "Sequential requests, no transactional snapshot. Late "
                "events/backfills and ranking changes can alter pages; "
                "duplicates/aggregate mismatches mark partial. Exhaustion is not "
                "proof of immutable completeness.",
            ),
            (
                "Daily distribution",
                "One whole-window server-grouped series, no adjacent query overlap. "
                "Server date buckets assumed UTC; deployed database timezone not "
                "verified. Missing days remain blank.",
            ),
            (
                "Index coverage",
                "One bounded request per clipped UTC day, inclusive start through "
                "clipped stop - 1s; "
                "fractional remainder excluded EVERY day to prevent overlapping "
                "buckets. crawl_date, not publication dates or distribution events. "
                "Current retained index, not historical inventory.",
            ),
            (
                "Index limits",
                "No per-article fanout. Index service may combine hot/archive stores "
                "for recent dates and mask backend failures; exact index completeness"
                " cannot be established.",
            ),
            (
                "Pagination",
                "Page size <=100; total_count is PAGE event sum, never full-domain "
                "total. Explicit page/next_page required; missing metadata, repeats, "
                "caps and failures mark partial.",
            ),
            (
                "Metric joining",
                "Stable article UUID, never sum duplicate rows or join on title. "
                "Missing metric is zero only after that ranking exhausts; otherwise "
                "blank. Metadata comes from named news.get_articles, not ranking fields. "
                "Missing enrichments never erase event counts or become zero.",
            ),
            (
                "Isolation",
                "One exact domain in every query; reject non-exact article "
                "hosts/share rows. No queries, consumer identities, billing data or "
                "raw response dumps exported.",
            ),
            (
                "SDK compatibility",
                "Verified against asknews 0.14.6; uses named sdk.news.get_articles for enrichment. "
                "Narrow authenticated client.request "
                "GET transport preserves grounded fields omitted by current typed "
                "distribution models; ranking/share/unique lack named wrappers. Index"
                " transport retains UTC offsets stripped by named method.",
            ),
            (
                "Permissions",
                "distribution for public rankings/share; internal + distribution for "
                "dashboard aggregates; news for enrichment/index. Admin status does not imply "
                "internal scope. No scope discovery or access changes.",
            ),
            (
                "No money",
                "Weighted share is not payout. No revenue pool, currency or dollars inferred.",
            ),
        ],
    )
    return book


def save_workbook(book, output):
    output = Path(output)
    if output.exists():
        raise ReportError("Output already exists; choose a new path (no overwrite).")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        os.link(
            temporary, output
        )  # Atomic publish, fails safely if destination appeared meanwhile.
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def credentials(environ):
    key = environ.get("ASKNEWS_API_KEY")
    client_id = environ.get("ASKNEWS_CLIENT_ID")
    secret = environ.get("ASKNEWS_CLIENT_SECRET")
    if key and (client_id or secret):
        raise ReportError("Choose API key OR OAuth environment credentials, not both.")
    if key:
        return {"api_key": key}
    if client_id and secret:
        scopes = set(environ.get("ASKNEWS_SCOPES", "distribution news").split())
        if not scopes or not scopes <= {"distribution", "news", "internal"}:
            raise ReportError(
                "ASKNEWS_SCOPES must list existing distribution/news/internal scopes."
            )
        return {"client_id": client_id, "client_secret": secret, "scopes": scopes}
    raise ReportError(
        "Set ASKNEWS_API_KEY or ASKNEWS_CLIENT_ID + ASKNEWS_CLIENT_SECRET "
        "securely in your local environment."
    )


def bounded_int(low, high):
    def convert(text):
        value = int(text)
        if not low <= value <= high:
            raise argparse.ArgumentTypeError(f"must be {low}–{high}")
        return value

    return convert


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--domain", default="apnews.com")
    parser.add_argument("--month", help="Optional YYYY-MM shorthand, exclusive with --start/--stop")
    parser.add_argument("--start", help="Inclusive timezone-aware ISO8601 datetime, whole seconds")
    parser.add_argument("--stop", help="Exclusive timezone-aware ISO8601 datetime, whole seconds")
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--include-internal",
        action="store_true",
        help="Try dashboard GETs only with existing authorized internal scope",
    )
    parser.add_argument(
        "--allow-partial",
        action="store_true",
        help="Save incomplete sections with explicit flags; exit 2",
    )
    parser.add_argument("--page-size", type=bounded_int(1, 100), default=100)
    parser.add_argument("--max-pages", type=bounded_int(1, 1000), default=100)
    parser.add_argument("--max-requests", type=bounded_int(1, 4000), default=400)
    parser.add_argument("--max-seconds", type=bounded_int(1, 3600), default=2400)
    args = parser.parse_args(argv)
    # SDK errors/HTTP logging can contain auth requests or response bodies. Never enable debug.
    logging.disable(logging.CRITICAL)
    try:
        domain = exact_domain(args.domain)
        window = select_window(args.month, args.start, args.stop)
        label = (
            window.start.strftime("%Y%m%dT%H%M%SZ") + "-" + window.stop.strftime("%Y%m%dT%H%M%SZ")
        )
        output = args.output or Path(f"{domain}-{label}-distribution.xlsx")
        if output.suffix.lower() != ".xlsx" or output.exists():
            raise ReportError("Choose a new .xlsx output path; existing files are not overwritten.")
        auth = credentials(os.environ)
        reader = Reader(None, args.max_requests, args.max_seconds)
        with AskNewsSDK(
            **auth,
            retries=0,
            timeout=30,
            follow_redirects=False,
            event_hooks={"request": [reader.before_request], "response": [reader.after_response]},
        ) as sdk:
            reader.sdk = sdk
            report = collect(
                reader, domain, window, args.page_size, args.max_pages, args.include_internal
            )
        report.coverage.append(
            ("HTTP sends", str(reader.requests), "Includes SDK-managed OAuth and bounded retries")
        )
        report.coverage.append(
            (
                "Configured limits",
                "bounded",
                f"page_size={args.page_size}; max_pages_per_metric={args.max_pages}; "
                f"max_http_sends={args.max_requests}; "
                f"request_start_budget_seconds={args.max_seconds}",
            )
        )
        if report.issues and not args.allow_partial:
            raise ReportError(
                "Incomplete report; no workbook written. Use --allow-partial to save "
                "flagged sections."
            )
        save_workbook(workbook(report), output)
        print(f"Saved {output}. {BOUNDARY}")
        return 2 if report.issues else 0
    except (ReportError, OSError) as exc:
        message = (
            str(exc)
            if isinstance(exc, ReportError)
            else "Cannot write workbook; check local path/permissions."
        )
        print(message, file=sys.stderr)
        return 1
    except Exception:
        print(
            "Unexpected SDK/schema failure; no report claimed complete. Verify "
            "pinned dependencies; no sensitive diagnostics emitted.",
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
